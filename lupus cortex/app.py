import io
import json
from pathlib import Path

import pandas as pd
import streamlit as st

st.set_page_config(page_title="Lupus Cortex | File → CSV", page_icon="◈", layout="wide")

st.markdown("""
<style>
.stApp{background:radial-gradient(circle at 15% 0%,rgba(91,140,255,.14),transparent 28%),radial-gradient(circle at 90% 5%,rgba(97,231,199,.10),transparent 26%),#070b12;color:#eef5ff}
.block-container{max-width:1180px;padding-top:2rem}
.brand{display:flex;align-items:center;gap:12px;margin-bottom:2.5rem}
.logo{width:44px;height:44px;border-radius:14px;display:grid;place-items:center;background:#10202a;border:1px solid rgba(97,231,199,.35);color:#61e7c7;font-weight:900}
.brand-title{font-size:18px;font-weight:800}.brand-sub{font-size:12px;color:#91a4bc}
.hero{text-align:center;margin:1rem auto 2rem;max-width:800px}.eyebrow{color:#61e7c7;font-size:11px;font-weight:800;letter-spacing:.14em;text-transform:uppercase}
.hero h1{font-size:clamp(2.4rem,6vw,4.5rem);line-height:1.02;margin:.7rem 0;color:#eef5ff}.hero p{color:#91a4bc;font-size:16px}
[data-testid="stFileUploader"]{background:linear-gradient(180deg,rgba(17,27,41,.95),rgba(9,16,26,.95));border:1px dashed #3b506a;border-radius:22px;padding:12px}
.stat-card{background:#0b131f;border:1px solid #203047;border-radius:15px;padding:16px}.stat-label{color:#71849b;font-size:10px;text-transform:uppercase;letter-spacing:.1em}.stat-value{font-size:20px;font-weight:800;margin-top:4px}
div.stButton>button,div.stDownloadButton>button{border-radius:11px;font-weight:750;min-height:44px}
div.stDownloadButton>button{background:linear-gradient(135deg,#61e7c7,#54bce7);color:#061016;border:0}
.notice{background:#121a25;border:1px solid #26364b;padding:12px 15px;border-radius:12px;color:#9db0c8;font-size:13px}
.footer{text-align:center;color:#60738b;font-size:12px;margin-top:30px}
</style>
""", unsafe_allow_html=True)

EXTENSIONS = {
    ".csv":"CSV",".tsv":"TSV",".xlsx":"XLSX",".xls":"XLS",".json":"JSON",
    ".xml":"XML",".txt":"TEXT",".log":"TEXT",".pdf":"PDF",".docx":"DOCX",
    ".html":"HTML",".htm":"HTML",".geojson":"GEOJSON",".kml":"KML",
    ".parquet":"PARQUET",".feather":"FEATHER"
}

def detect_type(f):
    ext = Path(f.name.lower()).suffix
    if ext in EXTENSIONS:
        return EXTENSIONS[ext]
    mime = (f.type or "").lower()
    if "spreadsheet" in mime or "excel" in mime: return "XLSX"
    if "json" in mime: return "JSON"
    if "pdf" in mime: return "PDF"
    if "word" in mime or "document" in mime: return "DOCX"
    if "html" in mime: return "HTML"
    if "text" in mime: return "TEXT"
    return "UNKNOWN"

def normalize(df):
    df = pd.DataFrame() if df is None else df.copy()
    if df.empty and len(df.columns) == 0:
        return pd.DataFrame({"content":[]})
    df = df.fillna("")
    cols = []
    seen = {}
    for i,c in enumerate(df.columns):
        name = str(c).strip()
        if not name or name.lower() in {"nan","none"}: name=f"column_{i+1}"
        seen[name] = seen.get(name,0)+1
        cols.append(name if seen[name]==1 else f"{name}_{seen[name]}")
    df.columns = cols
    return df

def json_to_df(data):
    if isinstance(data,list):
        if data and all(isinstance(x,dict) for x in data): return pd.json_normalize(data)
        return pd.DataFrame({"value":data})
    if isinstance(data,dict): return pd.DataFrame([data])
    return pd.DataFrame({"value":[data]})

def xml_to_df(raw):
    import xml.etree.ElementTree as ET
    root=ET.fromstring(raw); children=list(root)
    if not children: return pd.DataFrame({"content":[root.text or ""]})
    records=[]
    for child in children:
        r={"tag":child.tag}
        for sub in list(child): r[sub.tag]=(sub.text or "").strip()
        if len(r)==1: r["content"]=(child.text or "").strip()
        records.append(r)
    return pd.DataFrame(records)

def convert(f, kind):
    raw=f.getvalue()
    if kind=="CSV": return pd.read_csv(io.BytesIO(raw))
    if kind=="TSV": return pd.read_csv(io.BytesIO(raw),sep="\t")
    if kind in {"XLSX","XLS"}: return pd.read_excel(io.BytesIO(raw))
    if kind=="JSON": return json_to_df(json.loads(raw.decode("utf-8-sig")))
    if kind=="XML": return xml_to_df(raw)
    if kind=="TEXT": return pd.DataFrame({"content":[x for x in raw.decode("utf-8",errors="replace").splitlines() if x.strip()]})
    if kind=="HTML":
        tables=pd.read_html(io.StringIO(raw.decode("utf-8",errors="replace")))
        if tables: return tables[0]
        from bs4 import BeautifulSoup
        return pd.DataFrame({"content":[BeautifulSoup(raw,"html.parser").get_text(" ",strip=True)]})
    if kind=="PDF":
        import fitz
        doc=fitz.open(stream=raw,filetype="pdf")
        return pd.DataFrame([{"page":i,"content":p.get_text("text").strip()} for i,p in enumerate(doc,1)])
    if kind=="DOCX":
        from docx import Document
        doc=Document(io.BytesIO(raw))
        if doc.tables:
            rows=[[c.text.strip() for c in r.cells] for r in doc.tables[0].rows]
            if rows: return pd.DataFrame(rows[1:],columns=rows[0])
        return pd.DataFrame({"content":[p.text.strip() for p in doc.paragraphs if p.text.strip()]})
    if kind=="GEOJSON":
        data=json.loads(raw.decode("utf-8")); records=[]
        for feature in data.get("features",[]):
            r=dict(feature.get("properties") or {})
            r["geometry"]=json.dumps(feature.get("geometry"),ensure_ascii=False)
            records.append(r)
        return pd.DataFrame(records)
    if kind=="KML":
        return xml_to_df(raw)
    if kind=="PARQUET": return pd.read_parquet(io.BytesIO(raw))
    if kind=="FEATHER": return pd.read_feather(io.BytesIO(raw))
    raise ValueError("Unsupported file type")

st.markdown('<div class="brand"><div class="logo">LC</div><div><div class="brand-title">Lupus Cortex</div><div class="brand-sub">Universal Data Ingestion</div></div></div>',unsafe_allow_html=True)
st.markdown('<div class="hero"><div class="eyebrow">Urban Intelligence Data Pipeline</div><h1>Turn files into clean CSV data.</h1><p>Upload a file and Lupus Cortex detects its format, extracts structured data, previews the result, and prepares a downloadable CSV.</p></div>',unsafe_allow_html=True)

uploaded=st.file_uploader("Drop your file here or choose a file",type=None)

if uploaded:
    kind=detect_type(uploaded)
    cols=st.columns(4)
    values=[("File",uploaded.name),("Detected format",kind),("Size",f"{uploaded.size/1024:.1f} KB"),("Status","Ready")]
    for c,(label,value) in zip(cols,values):
        with c: st.markdown(f'<div class="stat-card"><div class="stat-label">{label}</div><div class="stat-value">{value}</div></div>',unsafe_allow_html=True)

    if st.button("Convert to CSV",type="primary",use_container_width=True):
        with st.spinner(f"Converting {kind} → CSV..."):
            try:
                df=normalize(convert(uploaded,kind))
                st.session_state["df"]=df
                st.session_state["source"]=uploaded.name
                st.success("Conversion completed successfully.")
            except Exception as e:
                st.error(f"Conversion failed: {e}")

if "df" in st.session_state:
    df=st.session_state["df"]; source=st.session_state["source"]
    st.markdown("### Conversion result")
    a,b,c=st.columns(3)
    a.metric("Rows",f"{len(df):,}"); b.metric("Columns",f"{len(df.columns):,}")
    csv=df.to_csv(index=False).encode("utf-8-sig")
    c.metric("CSV size",f"{len(csv)/1024:.1f} KB")
    st.markdown("### Preview")
    st.dataframe(df.head(100),use_container_width=True,height=380)
    st.download_button("Download CSV",data=csv,file_name=f"{Path(source).stem}_converted.csv",mime="text/csv",use_container_width=True)
    st.markdown('<div class="notice">Conversion is performed by the Python application. CSV output uses UTF-8 encoding for Bangla and Unicode data.</div>',unsafe_allow_html=True)

st.markdown('<div class="footer">Lupus Cortex · Universal File → CSV Converter · Python / Streamlit</div>',unsafe_allow_html=True)
